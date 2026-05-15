from collections import defaultdict
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


class DocGeneratorService:
    def generate_event_program(self, program_data: dict) -> bytes:
        doc = Document()

        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15

        sections = doc.sections[0]
        sections.top_margin = Cm(2)
        sections.bottom_margin = Cm(2)
        sections.left_margin = Cm(2.5)
        sections.right_margin = Cm(2)

        # Page header (optional: remove if not needed)
        header = sections.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = "Cogito, ergo sum. René Descartes"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.runs[0]
        run.font.size = Pt(9)
        run.font.italic = True

        event = program_data.get("event", {})
        committee = program_data.get("committee", [])
        all_sections = program_data.get("sections", [])

        self._build_title_page(doc, event)
        doc.add_page_break()

        self._build_committee(doc, event, committee)
        doc.add_page_break()

        self._build_schedule(doc, all_sections)
        doc.add_page_break()

        panel_sections = [s for s in all_sections if s.get("section_type") != "poster_session"]
        poster_sections = [s for s in all_sections if s.get("section_type") == "poster_session"]

        for section in panel_sections:
            self._build_section(doc, section)
            doc.add_page_break()

        if poster_sections:
            self._add_heading(doc, "POSTER SESSIONS", level=1)
            for section in poster_sections:
                self._build_section(doc, section)
                doc.add_page_break()

        self._build_russian_sections(doc, all_sections)

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _add_centered(self, doc, text, bold=True, size=14, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        return p

    def _add_heading(self, doc, text, level=1, centered=False):
        h = doc.add_heading(text, level=level)
        if centered:
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return h

    def _build_title_page(self, doc, event):
        self._add_centered(doc, "САРАТОВСКИЙ НАЦИОНАЛЬНЫЙ ИССЛЕДОВАТЕЛЬСКИЙ", size=14, space_after=4)
        self._add_centered(doc, "ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ ИМЕНИ Н.Г. ЧЕРНЫШЕВСКОГО", size=14, space_after=30)

        event_name = event.get("name", "")
        event_type = event.get("type", "")
        event_date = event.get("date", "")
        venue = event.get("venue", {})

        if event_type:
            self._add_centered(doc, event_type, size=13, space_after=8)
        self._add_centered(doc, f"«{event_name}»", size=16, space_after=16)

        if event_date:
            self._add_centered(doc, event_date, size=12, space_after=8)

        city = venue.get("city", "")
        if city:
            self._add_centered(doc, city, size=12, space_after=6)

    def _build_committee(self, doc, event, committee):
        if not committee:
            return

        event_name = event.get("name", "")
        self._add_heading(doc, f"Состав оргкомитета конференции «{event_name}»", level=1, centered=True)

        grouped = defaultdict(list)
        for m in committee:
            grouped[m.get("committee_type", "organizing")].append(m)

        for ctype_label, ctype_key in [("организационного", "organizing"), ("программного", "program")]:
            members = grouped.get(ctype_key, [])
            if not members:
                continue
            self._add_heading(doc, f"Члены {ctype_label} комитета:", level=2)

            for m in sorted(members, key=lambda x: x.get("sort_order", 0)):
                name = m.get("full_name", "")
                role = m.get("role", "")
                degree = m.get("degree") or m.get("title") or ""
                position = m.get("position") or ""
                workplace = m.get("workplace", "")

                parts = [name]
                if degree:
                    parts.append(f"({degree})")
                if position:
                    parts.append(f"— {position}")
                if workplace and workplace != position:
                    parts.append(f", {workplace}")

                line = " ".join(parts)
                if role:
                    line = f"{role}: {line}"

                p = doc.add_paragraph(line)
                p.paragraph_format.left_indent = Cm(1)

    def _build_schedule(self, doc, all_sections):
        if not all_sections:
            return

        self._add_heading(doc, "SCHEDULE", level=1, centered=True)

        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, label in enumerate(["Time", "Section", "Room"]):
            hdr[i].text = label
            for p in hdr[i].paragraphs:
                for run in p.runs:
                    run.bold = True

        for s in all_sections:
            time_str = s.get("time", "")
            time_display = ""
            try:
                dt = datetime.fromisoformat(time_str)
                time_display = dt.strftime("%H:%M")
            except (ValueError, TypeError):
                time_display = time_str
            name = s.get("name", "")
            hall = s.get("lecture_hall", "")
            row = table.add_row().cells
            row[0].text = time_display
            row[1].text = name
            row[2].text = hall

        doc.add_paragraph()

    def _build_section(self, doc, section):
        name = section.get("name", "")
        hall = section.get("lecture_hall", "")
        time_str = section.get("time", "")
        time_limit = section.get("time_limit")
        section_type = section.get("section_type", "")

        section_label = {
            "panel_discussion": "Panel Discussion",
            "poster_session": "Poster Session",
            "projects_session": "Projects Session",
        }.get(section_type, "Session")

        title_text = f"{section_label}: {name}"
        if hall:
            title_text += f" ({hall})"
        self._add_heading(doc, title_text, level=2)

        if time_str:
            try:
                dt = datetime.fromisoformat(time_str)
                date_display = dt.strftime("%d %B %Y")
            except (ValueError, TypeError):
                date_display = time_str
            p = doc.add_paragraph(date_display)
            p.paragraph_format.space_after = Pt(2)

        if time_limit:
            p = doc.add_paragraph(f"Time-limit: {time_limit} minutes")
            p.paragraph_format.space_after = Pt(4)

        juries = section.get("juries", [])
        if juries:
            p = doc.add_paragraph()
            run = p.add_run("Chairpersons:")
            run.bold = True
            for j in juries:
                jury_name = j.get("full_name", "")
                degree = j.get("degree") or j.get("title") or ""
                workplace = j.get("workplace") or ""
                university = j.get("university") or ""
                parts = [jury_name]
                if degree:
                    parts.append(f"({degree})")
                if workplace:
                    parts.append(workplace)
                if university and university != workplace:
                    parts.append(university)
                doc.add_paragraph(", ".join(parts))

        doc.add_paragraph()

        participants = section.get("participants", [])
        for idx, part in enumerate(participants, 1):
            name = part.get("full_name", "")
            topic = part.get("presentation_topic", "")
            abstract = part.get("abstract", "")
            advisor = part.get("scientific_advisor")

            p = doc.add_paragraph(f"{idx}. {name}")
            p.paragraph_format.space_after = Pt(2)

            if topic:
                p = doc.add_paragraph(topic)
                p.runs[0].italic = True
                p.paragraph_format.space_after = Pt(2)

            if abstract:
                p = doc.add_paragraph(abstract)
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.space_after = Pt(2)

            if advisor and advisor.get("full_name"):
                advisor_parts = [f"Scientific Advisor: {advisor['full_name']}"]
                adv_degree = advisor.get("degree") or advisor.get("title")
                if adv_degree:
                    advisor_parts.append(f"({adv_degree})")
                adv_workplace = advisor.get("workplace")
                if adv_workplace:
                    advisor_parts.append(adv_workplace)
                p = doc.add_paragraph(", ".join(advisor_parts))
                p.paragraph_format.space_after = Pt(6)

            if idx < len(participants):
                doc.add_paragraph()

    def _build_russian_sections(self, doc, all_sections):
        if not all_sections:
            return

        self._add_heading(doc, "СЕКЦИИ (русский)", level=1, centered=True)

        for section in all_sections:
            name = section.get("name", "")
            self._add_heading(doc, f"Секция: {name}", level=2)

            participants = section.get("participants", [])
            if not participants:
                doc.add_paragraph("Нет участников")
                continue

            for part in participants:
                name = part.get("full_name", "")
                faculty = part.get("faculty", "")
                university = part.get("university", "")
                topic = part.get("presentation_topic", "")

                parts = [name]
                place = ", ".join(filter(None, [university, faculty]))
                if place:
                    parts.append(f" ({place})")
                if topic:
                    parts.append(f" — {topic}")

                p = doc.add_paragraph(" ".join(parts))
                p.paragraph_format.left_indent = Cm(1)
